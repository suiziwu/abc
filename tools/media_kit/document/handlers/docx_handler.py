# -*- coding: utf-8 -*-
"""DOCX 处理器：段落 + 表格 + 标题层级 + 内嵌图片附件。"""

from __future__ import annotations

import os
import re
import zipfile
from pathlib import Path

from ...errors import ErrorCode, ParserError
from ...registry import register_document_handler
from ...types import AttachmentInfo, DocumentMetadata, DocumentPage, Heading
from .base import BaseDocumentHandler, DocumentExtract


def _heading_level(style_name: str) -> int:
    """把 Word 样式名映射为标题层级（中英文样式都支持）。"""
    name = (style_name or "").strip().lower()
    m = re.match(r"^(?:heading|标题)\s*([1-6])$", name)
    if m:
        return int(m.group(1))
    return 0


@register_document_handler(".docx")
class DocxHandler(BaseDocumentHandler):
    name = "docx"

    def extract(self, source: str, config, attachment_dir: str | None) -> DocumentExtract:
        try:
            from docx import Document
        except Exception as exc:  # pragma: no cover
            raise ParserError(
                ErrorCode.DEPENDENCY_MISSING, "缺少 python-docx：pip install python-docx", stage="dependency", cause=exc
            ) from exc

        try:
            doc = Document(source)
        except Exception as exc:
            raise ParserError(
                ErrorCode.DOCUMENT_TEXT_EXTRACT_FAILED, "DOCX 无法打开或已损坏", stage="extract", cause=exc
            ) from exc

        lines: list[str] = []
        structure: list[Heading] = []
        paragraphs = 0
        tables = 0

        def emit_paragraph(paragraph) -> None:
            nonlocal paragraphs
            text = (paragraph.text or "").strip()
            if not text:
                return
            paragraphs += 1
            level = _heading_level(getattr(paragraph.style, "name", ""))
            if level and config.structure.enabled and level <= config.structure.max_level:
                structure.append(Heading(level=level, title=text, page=None))
            lines.append(text)

        def emit_table(table) -> None:
            nonlocal tables
            tables += 1
            for row in table.rows:
                cells = [(cell.text or "").strip() for cell in row.cells]
                if any(cells):
                    lines.append(" | ".join(cells))

        try:
            if hasattr(doc, "iter_inner_content"):
                for block in doc.iter_inner_content():
                    if block.__class__.__name__ == "Table":
                        emit_table(block)
                    else:
                        emit_paragraph(block)
            else:  # 兼容旧版 python-docx
                for paragraph in doc.paragraphs:
                    emit_paragraph(paragraph)
                for table in doc.tables:
                    emit_table(table)
        except Exception as exc:
            raise ParserError(
                ErrorCode.DOCUMENT_TEXT_EXTRACT_FAILED, "读取 DOCX 正文失败", stage="extract", cause=exc
            ) from exc

        props = doc.core_properties
        metadata = DocumentMetadata(
            path=str(source),
            format=Path(source).suffix.lower(),
            size_bytes=os.path.getsize(source),
            page_count=0,
            title=(props.title or "") or "",
            author=(props.author or "") or "",
            extra={"paragraphs": paragraphs, "tables": tables},
        )
        page = DocumentPage(index=1, text="\n".join(lines))
        attachments = self._attachments(source, attachment_dir, config) if attachment_dir else []

        return DocumentExtract(
            metadata=metadata,
            pages=[page],
            structure=structure,
            attachments=attachments,
        )

    @staticmethod
    def _attachments(source: str, attachment_dir: str, config) -> list[AttachmentInfo]:
        result: list[AttachmentInfo] = []
        try:
            with zipfile.ZipFile(source) as zf:
                for name in zf.namelist():
                    if not name.startswith("word/media/") or name.endswith("/"):
                        continue
                    data = zf.read(name)
                    if len(data) < config.attachments.min_size_bytes:
                        continue
                    filename = os.path.basename(name)
                    path = os.path.join(attachment_dir, filename)
                    with open(path, "wb") as fp:
                        fp.write(data)
                    result.append(
                        AttachmentInfo(
                            name=filename,
                            path=path,
                            content_type=f"image/{filename.rsplit('.', 1)[-1].lower()}",
                            source=f"docx:{name}",
                            size_bytes=len(data),
                        )
                    )
        except Exception:
            return result
        return result
