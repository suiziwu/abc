# -*- coding: utf-8 -*-
"""表格处理器：xlsx / xlsm / csv，按工作表输出文本与结构。"""

from __future__ import annotations

import csv as csv_module
import io
import os
from pathlib import Path

from ...errors import ErrorCode, ParserError
from ...registry import register_document_handler
from ...types import DocumentMetadata, DocumentPage, Heading
from .base import BaseDocumentHandler, DocumentExtract


def _read_csv_text(source: str) -> str:
    raw = Path(source).read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gbk", "latin-1"):
        try:
            text = raw.decode(encoding)
        except UnicodeDecodeError:
            continue
        try:
            dialect = csv_module.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv_module.Error:
            dialect = csv_module.excel
        reader = csv_module.reader(io.StringIO(text), dialect)
        return "\n".join("\t".join(cell.strip() for cell in row) for row in reader if any(row))
    raise ParserError(ErrorCode.DOCUMENT_TEXT_EXTRACT_FAILED, "CSV 编码无法识别", stage="extract")


@register_document_handler(".xlsx", ".xlsm", ".csv")
class SpreadsheetHandler(BaseDocumentHandler):
    name = "spreadsheet"

    def extract(self, source: str, config, attachment_dir: str | None) -> DocumentExtract:
        suffix = Path(source).suffix.lower()
        if suffix == ".csv":
            text = _read_csv_text(source)
            metadata = DocumentMetadata(
                path=str(source),
                format=suffix,
                size_bytes=os.path.getsize(source),
                page_count=1,
                extra={"sheets": 1},
            )
            return DocumentExtract(metadata=metadata, pages=[DocumentPage(index=1, text=text)], structure=[])

        try:
            from openpyxl import load_workbook
        except Exception as exc:  # pragma: no cover
            raise ParserError(
                ErrorCode.DEPENDENCY_MISSING, "缺少 openpyxl：pip install openpyxl", stage="dependency", cause=exc
            ) from exc

        try:
            wb = load_workbook(source, read_only=True, data_only=True)
        except Exception as exc:
            raise ParserError(
                ErrorCode.DOCUMENT_TEXT_EXTRACT_FAILED, "表格无法打开或已损坏", stage="extract", cause=exc
            ) from exc

        pages: list[DocumentPage] = []
        structure: list[Heading] = []
        try:
            for index, sheet in enumerate(wb.worksheets, 1):
                lines: list[str] = []
                for row in sheet.iter_rows(values_only=True):
                    cells = ["" if value is None else str(value) for value in row]
                    while cells and not cells[-1]:
                        cells.pop()
                    if any(cells):
                        lines.append("\t".join(cells))
                pages.append(DocumentPage(index=index, text="\n".join(lines), meta={"sheet": sheet.title}))
                if config.structure.enabled:
                    structure.append(Heading(level=1, title=sheet.title, page=index))
            props = wb.properties
            metadata = DocumentMetadata(
                path=str(source),
                format=suffix,
                size_bytes=os.path.getsize(source),
                page_count=len(pages),
                title=(props.title or "") if props else "",
                author=(props.creator or "") if props else "",
                extra={"sheets": [sheet.title for sheet in wb.worksheets]},
            )
        finally:
            wb.close()

        return DocumentExtract(metadata=metadata, pages=pages, structure=structure)
