# -*- coding: utf-8 -*-
"""PPTX 处理器：按幻灯片提取文本、标题结构与内嵌媒体。"""

from __future__ import annotations

import os
import zipfile
from pathlib import Path

from ...errors import ErrorCode, ParserError
from ...registry import register_document_handler
from ...types import AttachmentInfo, DocumentMetadata, DocumentPage, Heading
from .base import BaseDocumentHandler, DocumentExtract


def _shape_texts(shape) -> list[str]:
    texts: list[str] = []
    if getattr(shape, "has_text_frame", False) and shape.text_frame is not None:
        text = (shape.text_frame.text or "").strip()
        if text:
            texts.append(text)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [(cell.text or "").strip() for cell in row.cells]
            if any(cells):
                texts.append(" | ".join(cells))
    if getattr(shape, "shapes", None):  # 组合图形
        for child in shape.shapes:
            texts.extend(_shape_texts(child))
    return texts


@register_document_handler(".pptx")
class PptxHandler(BaseDocumentHandler):
    name = "pptx"

    def extract(self, source: str, config, attachment_dir: str | None) -> DocumentExtract:
        try:
            from pptx import Presentation
        except Exception as exc:  # pragma: no cover
            raise ParserError(
                ErrorCode.DEPENDENCY_MISSING, "缺少 python-pptx：pip install python-pptx", stage="dependency", cause=exc
            ) from exc

        try:
            prs = Presentation(source)
        except Exception as exc:
            raise ParserError(
                ErrorCode.DOCUMENT_TEXT_EXTRACT_FAILED, "PPTX 无法打开或已损坏", stage="extract", cause=exc
            ) from exc

        pages: list[DocumentPage] = []
        structure: list[Heading] = []
        for index, slide in enumerate(prs.slides, 1):
            texts: list[str] = []
            for shape in slide.shapes:
                texts.extend(_shape_texts(shape))
            body = "\n".join(texts)
            pages.append(DocumentPage(index=index, text=body, meta={"slide": index}))
            if config.structure.enabled and texts:
                title = texts[0].splitlines()[0][:60]
                structure.append(Heading(level=1, title=f"第{index}页 {title}", page=index))

        props = prs.core_properties
        metadata = DocumentMetadata(
            path=str(source),
            format=Path(source).suffix.lower(),
            size_bytes=os.path.getsize(source),
            page_count=len(prs.slides),
            title=(props.title or "") or "",
            author=(props.author or "") or "",
            extra={"slides": len(prs.slides)},
        )
        attachments = self._attachments(source, attachment_dir, config) if attachment_dir else []

        return DocumentExtract(
            metadata=metadata,
            pages=pages,
            structure=structure,
            attachments=attachments,
        )

    @staticmethod
    def _attachments(source: str, attachment_dir: str, config) -> list[AttachmentInfo]:
        result: list[AttachmentInfo] = []
        try:
            with zipfile.ZipFile(source) as zf:
                for name in zf.namelist():
                    if not name.startswith("ppt/media/") or name.endswith("/"):
                        continue
                    data = zf.read(name)
                    if len(data) < config.attachments.min_size_bytes:
                        continue
                    filename = os.path.basename(name)
                    path = os.path.join(attachment_dir, filename)
                    with open(path, "wb") as fp:
                        fp.write(data)
                    result.append(
                        AttachmentInfo(
                            name=filename,
                            path=path,
                            content_type=f"image/{filename.rsplit('.', 1)[-1].lower()}",
                            source=f"pptx:{name}",
                            size_bytes=len(data),
                        )
                    )
        except Exception:
            return result
        return result
