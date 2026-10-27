# -*- coding: utf-8 -*-
"""测试夹具生成：用 ffmpeg / PyMuPDF / python-docx 等现场生成小样本。"""

from __future__ import annotations

import os
import shutil
import subprocess
from pathlib import Path

FONT_CANDIDATES = [
    r"C:\Windows\Fonts\simhei.ttf",
    r"C:\Windows\Fonts\msyh.ttc",
    r"C:\Windows\Fonts\simsun.ttc",
    "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
]
FONT = next((f for f in FONT_CANDIDATES if os.path.exists(f)), None)


def ffmpeg_exe() -> str:
    import imageio_ffmpeg

    return imageio_ffmpeg.get_ffmpeg_exe()


def make_video(path: str, *, duration: float = 1.0) -> str:
    """生成 320x240、带 440Hz 正弦音轨的短视频。"""
    ext = Path(path).suffix.lower()
    cmd = [
        ffmpeg_exe(), "-y", "-hide_banner", "-loglevel", "error",
        "-f", "lavfi", "-i", f"color=c=navy:s=320x240:d={duration}",
        "-f", "lavfi", "-i", f"sine=frequency=440:duration={duration}",
    ]
    if ext == ".webm":
        cmd += ["-c:v", "libvpx", "-b:v", "200k", "-c:a", "libopus", "-shortest"]
    else:
        cmd += ["-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", "-shortest"]
    if ext == ".ts":
        cmd += ["-f", "mpegts"]
    elif ext == ".mkv":
        cmd += ["-f", "matroska"]
    cmd.append(path)
    subprocess.run(cmd, check=True)
    return path


def make_pdf(path: str, *, with_toc: bool = False) -> str:
    """生成两页中文 PDF；可选写入书签。"""
    import pymupdf

    doc = pymupdf.open()
    lines = [
        ("第1章 计算机网络", 20, 72, 1),
        ("1.1 计算机网络的定义", 15, 72, 110),
        ("计算机网络是由通信信道连接的主机和网络设备的集合。", 12, 72, 150),
        ("1.1.1 定义", 14, 72, 200),
        ("计算机网络最主要的特征是通用性。", 12, 72, 240),
    ]
    for page_index in range(2):
        page = doc.new_page()
        for text, size, x, y in lines:
            if FONT:
                page.insert_text((x, y), text, fontsize=size, fontname="china", fontfile=FONT)
            else:
                page.insert_text((x, y), f"page{page_index + 1} " + text.encode("ascii", "ignore").decode(), fontsize=size)
        page.insert_text((72, 300), f"第 {page_index + 1} 页内容。", fontsize=11, fontname="china", fontfile=FONT) if FONT \
            else page.insert_text((72, 300), f"page {page_index + 1}", fontsize=11)
    if with_toc:
        doc.set_toc([[1, "第1章 计算机网络", 1], [2, "1.1 计算机网络的定义", 1], [2, "1.2 因特网的结构", 2]])
    doc.save(path)
    doc.close()
    return path


def make_scanned_pdf(path: str) -> str:
    """把中文页面渲染成图片再包进 PDF，模拟扫描件。"""
    import pymupdf

    src = pymupdf.open()
    page = src.new_page()
    if FONT:
        page.insert_text((72, 100), "计算机网络的定义", fontsize=28, fontname="china", fontfile=FONT)
        page.insert_text((72, 160), "计算机网络是由主机和网络设备组成的系统。", fontsize=16, fontname="china", fontfile=FONT)
    else:
        page.insert_text((72, 100), "computer network", fontsize=28)
    pix = page.get_pixmap(dpi=150)
    png = pix.tobytes("png")
    out = pymupdf.open()
    new_page = out.new_page(width=page.rect.width, height=page.rect.height)
    new_page.insert_image(new_page.rect, stream=png)
    out.save(path)
    out.close()
    src.close()
    return path


def make_docx(path: str) -> str:
    from docx import Document
    from docx.shared import Inches
    from PIL import Image

    img_path = str(Path(path).with_suffix(".png"))
    Image.new("RGB", (40, 40), (30, 120, 200)).save(img_path)

    doc = Document()
    doc.add_heading("计算机网络", level=1)
    doc.add_heading("网络协议", level=2)
    doc.add_paragraph("协议是通信双方必须遵守的规则。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "层次"
    table.cell(0, 1).text = "功能"
    table.cell(1, 0).text = "运输层"
    table.cell(1, 1).text = "端到端传输"
    doc.add_picture(img_path, width=Inches(0.5))
    doc.save(path)
    return path


def make_pptx(path: str) -> str:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "计算机网络"
    slide.placeholders[1].text = "第一章 概述"
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "网络协议"
    slide2.placeholders[1].text = "协议与分层"
    prs.save(path)
    return path


def make_xlsx(path: str) -> str:
    from openpyxl import Workbook

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "基础"
    ws1.append(["名称", "数值"])
    ws1.append(["时延", 100])
    ws2 = wb.create_sheet("进阶")
    ws2.append(["协议", "层次"])
    ws2.append(["TCP", "运输层"])
    wb.save(path)
    return path


def make_csv(path: str) -> str:
    Path(path).write_text("名称,数值\n时延,100\n带宽,200\n", encoding="utf-8")
    return path


def make_md(path: str) -> str:
    Path(path).write_text(
        "# 计算机网络\n\n## 1.1 网络定义\n\n计算机网络是互联的自治系统集合。\n",
        encoding="utf-8",
    )
    return path


def make_txt(path: str) -> str:
    Path(path).write_text("第1章 计算机网络\n1.1 网络定义\n正文内容。\n", encoding="utf-8")
    return path


def make_html(path: str) -> str:
    Path(path).write_text(
        "<html><body><h1>计算机网络</h1><p>网络协议是通信规则。</p>"
        "<h2>1.1 网络定义</h2><p>主机与链路。</p><script>hidden()</script></body></html>",
        encoding="utf-8",
    )
    return path


def cleanup(path: str) -> None:
    shutil.rmtree(path, ignore_errors=True)
