import os
import pandas as pd
import unicodedata
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ===================== 路径配置（统一指向工程根目录的 04-清洗后数据）=====================
_BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_DATA_DIR = os.path.join(_BASE_DIR, "04-清洗后数据")
os.makedirs(_DATA_DIR, exist_ok=True)
INPUT_CSV = os.path.join(_DATA_DIR, "AI全栈岗位_去空技能.csv")
OUTPUT_CSV = os.path.join(_DATA_DIR, "AI全栈开发工程师岗位_清洗后.csv")
OUTPUT_EXCEL = os.path.join(_DATA_DIR, "AI全栈开发工程师岗位_清洗后.xlsx")
JOB_KEY = "AI全栈开发工程师"
# ======================================================================

# 读取原始数据
df = pd.read_csv(INPUT_CSV)

# 筛选岗位：仅保留包含关键词的行
df_filtered = df[df["职位名称"].str.contains(JOB_KEY, na=False)].copy()

# 重置序号列
df_filtered["序号"] = list(range(1, len(df_filtered) + 1))

# 保存CSV
df_filtered.to_csv(OUTPUT_CSV, index=False, encoding="utf-8-sig")
print(f"✅ CSV 保存成功：{OUTPUT_CSV}")

# 生成Excel文件
with pd.ExcelWriter(OUTPUT_EXCEL, engine="openpyxl") as writer:
    df_filtered.to_excel(writer, sheet_name="AI全栈开发工程师岗位", index=False)

# Excel样式美化
wb = load_workbook(OUTPUT_EXCEL)
ws = wb.active

# 样式定义
header_fill = PatternFill("solid", fgColor="0070C0")
header_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
data_font = Font(name="微软雅黑", size=10, color="333333")
center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
thin_border = Side(style="thin", color="D9D9D9")
border = Border(bottom=thin_border)
zebra_blue = "EBF1F8"
white = "FFFFFF"

# 表头样式
for col in range(1, ws.max_column + 1):
    cell = ws.cell(row=1, column=col)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center_align
    cell.border = border

# 数据行斑马纹样式
for row in range(2, ws.max_row + 1):
    bg_color = zebra_blue if (row - 2) % 2 == 0 else white
    row_fill = PatternFill("solid", fgColor=bg_color)
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = data_font
        cell.fill = row_fill
        cell.border = border
        cell.alignment = center_align if col == 1 else left_align

# 自动适配列宽
def display_width(text):
    return sum(2 if unicodedata.east_asian_width(c) in ("F", "W") else 1 for c in str(text or ""))

def auto_fit_columns(ws, min_w=8, max_w=50, padding=3):
    for col in ws.columns:
        col_letter = col[0].column_letter
        max_len = 0
        for cell in col:
            if cell.value is not None and not isinstance(cell, type(ws.merged_cells)):
                current_len = display_width(cell.value)
                if current_len > max_len:
                    max_len = current_len
        col_width = min(max(max_len * 1.1 + padding, min_w), max_w)
        ws.column_dimensions[col_letter].width = col_width

auto_fit_columns(ws)

# 冻结表头
ws.freeze_panes = "A2"

# 保存Excel
wb.save(OUTPUT_EXCEL)
print(f"✅ Excel 保存成功：{OUTPUT_EXCEL}")
print(f"📊 筛选完成，共保留 {len(df_filtered)} 条 AI全栈开发工程师 岗位数据")